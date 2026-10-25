"""Import d'entrepôts depuis un CSV ou un Excel. Rien n'est enregistré si une ligne est invalide."""

import csv
import io
import re
import unicodedata
from datetime import datetime, time
from decimal import Decimal, InvalidOperation

from django.core.exceptions import ValidationError
from django.db import IntegrityError, transaction
from openpyxl import Workbook, load_workbook

from apps.shipments.forms import GOUVERNORATS
from apps.warehouses.models import Entrepot

MAX_OCTETS = 2 * 1024 * 1024
MAX_LIGNES = 500
COLONNES = (
    "nom",
    "adresse",
    "ville",
    "gouvernorat",
    "code_postal",
    "latitude",
    "longitude",
    "capacite_max_m3",
    "horaire_ouverture",
    "horaire_fermeture",
    "duree_stockage_max_heures",
    "delai_grace_minutes",
    "actif",
    "description",
)
OBLIGATOIRES = ("nom", "adresse", "ville", "capacite_max_m3")
ALIAS = {
    "nom": "nom",
    "name": "nom",
    "entrepot": "nom",
    "adresse": "adresse",
    "address": "adresse",
    "ville": "ville",
    "city": "ville",
    "gouvernorat": "gouvernorat",
    "code_postal": "code_postal",
    "codepostal": "code_postal",
    "latitude": "latitude",
    "lat": "latitude",
    "longitude": "longitude",
    "lng": "longitude",
    "lon": "longitude",
    "capacite_max_m3": "capacite_max_m3",
    "capacite": "capacite_max_m3",
    "capacite_m3": "capacite_max_m3",
    "capacite_maximale_m3": "capacite_max_m3",
    "horaire_ouverture": "horaire_ouverture",
    "ouverture": "horaire_ouverture",
    "horaire_d_ouverture": "horaire_ouverture",
    "horaire_fermeture": "horaire_fermeture",
    "fermeture": "horaire_fermeture",
    "horaire_de_fermeture": "horaire_fermeture",
    "duree_stockage_max_heures": "duree_stockage_max_heures",
    "duree_stockage": "duree_stockage_max_heures",
    "duree_stockage_heures": "duree_stockage_max_heures",
    "delai_grace_minutes": "delai_grace_minutes",
    "delai_grace": "delai_grace_minutes",
    "actif": "actif",
    "description": "description",
}
EXEMPLE = {
    "nom": "Relais Exemple",
    "adresse": "Avenue Habib Bourguiba",
    "ville": "Tunis",
    "gouvernorat": "Tunis",
    "code_postal": "1000",
    "latitude": "36.8065",
    "longitude": "10.1815",
    "capacite_max_m3": "25",
    "horaire_ouverture": "08:00",
    "horaire_fermeture": "18:00",
    "duree_stockage_max_heures": "48",
    "delai_grace_minutes": "30",
    "actif": "oui",
    "description": "Supprimez cette ligne d'exemple avant un import réel.",
}


def contenu_csv():
    sortie = io.StringIO()
    writer = csv.DictWriter(sortie, fieldnames=COLONNES, lineterminator="\n")
    writer.writeheader()
    writer.writerow(EXEMPLE)
    return ("\ufeff" + sortie.getvalue()).encode()


def contenu_xlsx():
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Entrepots"
    feuille.append(list(COLONNES))
    feuille.append([EXEMPLE[colonne] for colonne in COLONNES])
    sortie = io.BytesIO()
    classeur.save(sortie)
    return sortie.getvalue()


def importer(fichier):
    """Retourne (nombre créé, erreurs). En cas d'erreur, aucun entrepôt n'est créé."""
    try:
        tableau = _lire(fichier)
    except ValueError as exc:
        return 0, [str(exc)]
    if isinstance(tableau, str):
        return 0, [tableau]
    prets, erreurs = _valider(tableau)
    if erreurs:
        return 0, erreurs
    try:
        with transaction.atomic():
            for entrepot in prets:
                entrepot.save()
    except IntegrityError:
        return 0, ["Un entrepôt porte déjà ce nom dans cette ville. Aucun entrepôt n'a été créé."]
    return len(prets), []


def _lire(fichier):
    nom = (getattr(fichier, "name", "") or "").lower()
    if getattr(fichier, "size", 0) > MAX_OCTETS:
        raise ValueError("Le fichier dépasse 2 Mo.")
    contenu = fichier.read()
    if nom.endswith(".csv"):
        return _csv(contenu)
    if nom.endswith(".xlsx"):
        return _xlsx(contenu)
    raise ValueError("Utilisez un fichier .csv ou .xlsx.")


def _csv(contenu):
    try:
        texte = contenu.decode("utf-8-sig")
    except UnicodeDecodeError:
        texte = contenu.decode("latin-1")
    delimiteur = ";" if texte[:2000].count(";") > texte[:2000].count(",") else ","
    lecteur = csv.DictReader(io.StringIO(texte), delimiter=delimiteur)
    if not lecteur.fieldnames:
        return "Le fichier ne contient pas d'en-tête."
    return _lignes(lecteur.fieldnames, lecteur)


def _xlsx(contenu):
    try:
        classeur = load_workbook(io.BytesIO(contenu), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("Le fichier Excel est illisible.") from exc
    feuille = classeur.active
    lignes = feuille.iter_rows(values_only=True)
    try:
        entetes = next(lignes)
    except StopIteration:
        return "Le fichier ne contient pas d'en-tête."
    return _lignes(entetes, lignes)


def _lignes(entetes, lignes):
    colonnes = [_entete(valeur) for valeur in entetes]
    manquantes = [nom for nom in OBLIGATOIRES if nom not in colonnes]
    if manquantes:
        return "Colonnes obligatoires manquantes : " + ", ".join(manquantes) + "."
    resultat = []
    for indice, ligne in enumerate(lignes, start=2):
        if hasattr(ligne, "keys"):
            valeurs = list(ligne.values())
        else:
            valeurs = list(ligne)
        if all(_vide(valeur) for valeur in valeurs):
            continue
        resultat.append((indice, {colonnes[i]: valeurs[i] for i in range(min(len(colonnes), len(valeurs)))}))
        if len(resultat) > MAX_LIGNES:
            return f"Le fichier dépasse {MAX_LIGNES} entrepôts."
    if not resultat:
        return "Le fichier ne contient aucune ligne d'entrepôt."
    return resultat


def _valider(tableau):
    existants = set(Entrepot.objects.values_list("nom", "ville"))
    vus = set()
    prets = []
    erreurs = []
    gouvernorats = {_clef(nom): nom for nom, _libelle in GOUVERNORATS}
    for numero, ligne in tableau:
        entrepot, erreur = _entrepot(ligne, gouvernorats)
        if erreur:
            erreurs.append(f"Ligne {numero} : {erreur}")
            continue
        cle = (entrepot.nom, entrepot.ville)
        if cle in vus:
            erreurs.append(f"Ligne {numero} : ce nom est déjà présent dans le fichier pour {entrepot.ville}.")
            continue
        if cle in existants:
            erreurs.append(f"Ligne {numero} : un entrepôt porte déjà ce nom dans {entrepot.ville}.")
            continue
        try:
            entrepot.full_clean()
        except ValidationError as exc:
            erreurs.append(f"Ligne {numero} : {' '.join(exc.messages)}")
            continue
        vus.add(cle)
        prets.append(entrepot)
    return prets, erreurs


def _entrepot(ligne, gouvernorats):
    try:
        nom = _borne(ligne.get("nom"), 120, "nom")
        adresse = _borne(ligne.get("adresse"), 255, "adresse")
        ville = _borne(ligne.get("ville"), 80, "ville")
        if not nom or not adresse or not ville:
            return None, "le nom, l'adresse et la ville sont obligatoires."
        capacite = _decimal(ligne.get("capacite_max_m3"))
        if capacite is None or capacite <= 0:
            return None, "la capacité doit être strictement positive."
        gouvernorat = _borne(ligne.get("gouvernorat"), 40, "gouvernorat")
        if gouvernorat:
            connu = gouvernorats.get(_clef(gouvernorat))
            if connu is None:
                return None, f"le gouvernorat « {gouvernorat} » n'est pas reconnu."
            gouvernorat = connu
        latitude = _flottant(ligne.get("latitude"), -90, 90, "latitude")
        longitude = _flottant(ligne.get("longitude"), -180, 180, "longitude")
        ouverture = _heure(ligne.get("horaire_ouverture"), time(8, 0), "ouverture")
        fermeture = _heure(ligne.get("horaire_fermeture"), time(18, 0), "fermeture")
        duree = _entier(ligne.get("duree_stockage_max_heures"), 48, 1, "durée de stockage")
        grace = _entier(ligne.get("delai_grace_minutes"), 30, 0, "délai de grâce")
        actif = _booleen(ligne.get("actif"))
        description = _borne(ligne.get("description"), 2000, "description")
        code_postal = _borne(ligne.get("code_postal"), 10, "code postal")
    except ValueError as exc:
        return None, str(exc)
    return (
        Entrepot(
            nom=nom,
            description=description,
            adresse=adresse,
            ville=ville,
            gouvernorat=gouvernorat,
            code_postal=code_postal,
            latitude=latitude,
            longitude=longitude,
            capacite_max_m3=capacite,
            horaire_ouverture=ouverture,
            horaire_fermeture=fermeture,
            duree_stockage_max_heures=duree,
            delai_grace_minutes=grace,
            actif=actif,
        ),
        "",
    )


def _entete(valeur):
    return ALIAS.get(_clef(valeur), _clef(valeur))


def _clef(valeur):
    texte = unicodedata.normalize("NFKD", str(valeur or ""))
    texte = "".join(caractere for caractere in texte if not unicodedata.combining(caractere))
    return re.sub(r"[^a-z0-9]+", "_", texte.lower()).strip("_")


def _vide(valeur):
    return valeur is None or str(valeur).strip() == ""


def _borne(valeur, limite, libelle):
    if _vide(valeur):
        return ""
    texte = str(valeur).strip()
    if len(texte) > limite:
        raise ValueError(f"le champ {libelle} dépasse {limite} caractères.")
    return texte


def _decimal(valeur):
    if _vide(valeur):
        return None
    texte = str(valeur).strip().replace(" ", "").replace(",", ".")
    try:
        return Decimal(texte)
    except InvalidOperation as exc:
        raise ValueError("la capacité n'est pas un nombre.") from exc


def _flottant(valeur, minimum, maximum, libelle):
    if _vide(valeur):
        return None
    try:
        nombre = float(str(valeur).strip().replace(",", "."))
    except ValueError as exc:
        raise ValueError(f"la {libelle} n'est pas un nombre.") from exc
    if not minimum <= nombre <= maximum:
        raise ValueError(f"la {libelle} est hors limites.")
    return nombre


def _heure(valeur, defaut, libelle):
    if _vide(valeur):
        return defaut
    if isinstance(valeur, datetime):
        return valeur.time().replace(microsecond=0)
    if isinstance(valeur, time):
        return valeur.replace(microsecond=0)
    texte = str(valeur).strip()
    for motif in ("%H:%M", "%H:%M:%S"):
        try:
            return datetime.strptime(texte, motif).time()
        except ValueError:
            continue
    raise ValueError(f"l'heure de {libelle} doit être au format HH:MM.")


def _entier(valeur, defaut, minimum, libelle):
    if _vide(valeur):
        return defaut
    try:
        nombre = int(Decimal(str(valeur).strip().replace(",", ".")))
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"la {libelle} doit être un nombre entier.") from exc
    if nombre < minimum:
        raise ValueError(f"la {libelle} doit être au moins {minimum}.")
    return nombre


def _booleen(valeur):
    if _vide(valeur):
        return True
    if isinstance(valeur, bool):
        return valeur
    texte = _clef(valeur)
    if texte in {"oui", "yes", "true", "1", "actif", "vrai"}:
        return True
    if texte in {"non", "no", "false", "0", "inactif", "faux"}:
        return False
    raise ValueError("la colonne actif doit contenir oui ou non.")
